"""
Veille CNIL, deux sources :
 1) https://www.cnil.fr/fr/actualite : articles portant le tag #Sanction
    (ligne détaillée : entité, articles du RGPD, lien vers l'article) ;
 2) https://www.cnil.fr/fr/les-sanctions-prononcees-par-la-cnil : le tableau
    officiel de TOUTES les sanctions prononcées depuis 2011, y compris celles
    rendues en procédure simplifiée. Au premier passage, tout l'historique est
    repris ; ensuite, seules les lignes nouvelles sont ajoutées.
Les doublons sont évités grâce au lien Légifrance (colonne H), au lien de
l'article (colonne G) et à une clé technique (colonne J, masquée).

Usage :  python veille_cnil.py            (met à jour le fichier)
         python veille_cnil.py --dry-run  (affiche sans rien écrire)
"""
import re
import sys
import datetime as dt
from pathlib import Path
from urllib.parse import urljoin

import html as htmlmod
import shutil
import time
from zoneinfo import ZoneInfo

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from copy import copy

# ---------------- Réglages ----------------
FICHIER_EXCEL = Path(__file__).with_name("Sanctions_CNIL.xlsx")
DOSSIER_SITE = Path(__file__).with_name("site")  # page web publiée par GitHub Pages
URL_ACTUS = "https://www.cnil.fr/fr/actualite"
NB_PAGES = 3                      # pages de la liste à parcourir (~6 articles/page)
TITRES_EXCLUS = ["clôture de l’injonction", "clôture de l'injonction"]  # tag #Sanction mais pas une sanction
SURLIGNER_NOUVEAU = PatternFill("solid", fgColor="FFF2CC")  # jaune pâle = ligne à relire
COL_LIEN, COL_LEGI, COL_AJOUT = 7, 8, 9   # G = lien article CNIL, H = lien délibération Légifrance, I = date d'ajout automatique
COL_CLE = 10                              # J = clé technique de la ligne du tableau officiel CNIL (masquée)
URL_TABLEAU = "https://www.cnil.fr/fr/les-sanctions-prononcees-par-la-cnil"
MAX_LEGIFRANCE = 250                      # nb max de délibérations Légifrance lues par passage (pour retrouver le nom de l'entité)
LIBELLE_SIMPLIFIEE = "Procédure simplifiée"
HEADERS = {"User-Agent": "Mozilla/5.0 (veille sanctions CNIL, usage personnel)"}
# ------------------------------------------

MOIS = {m: i for i, m in enumerate(
    ["janvier", "février", "mars", "avril", "mai", "juin", "juillet",
     "août", "septembre", "octobre", "novembre", "décembre"], 1)}
RE_DATE = re.compile(r"(\d{1,2})(?:er)?\s+(" + "|".join(MOIS) + r")\s+(\d{4})", re.I)


def parse_date(txt):
    m = RE_DATE.search(txt or "")
    if not m:
        return None
    return dt.datetime(int(m.group(3)), MOIS[m.group(2).lower()], int(m.group(1)))


def telecharger(url):
    r = requests.get(url, headers=HEADERS, timeout=30)
    r.raise_for_status()
    return r.text


def extraire_articles(html, base=URL_ACTUS):
    """Renvoie la liste des articles de la page : titre, lien, résumé, tags, date de publication."""
    soup = BeautifulSoup(html, "html.parser")
    articles = {}
    for h3 in soup.find_all("h3"):
        a = h3.find("a", href=True)
        if not a:
            continue
        lien = urljoin(base, a["href"])
        # remonte jusqu'au bloc qui contient CET article uniquement (et sa date)
        bloc = h3
        for _ in range(8):
            parent = bloc.parent
            if parent is None:
                break
            liens_h3 = {urljoin(base, x["href"]) for h in parent.find_all("h3") for x in h.find_all("a", href=True)}
            if len(liens_h3) > 1:
                break
            bloc = parent
        texte = bloc.get_text(" ", strip=True)
        tags = [t.get_text(strip=True) for t in bloc.find_all("a", href=re.compile(r"/tag/"))]
        if not tags:
            tags = re.findall(r"#[\wÀ-ÿ’' -]+?(?=\s#|\s\d|$)", texte)
        titre = a.get_text(" ", strip=True)
        resume = texte.replace(titre, "", 1)
        resume = re.split(r"\s#", resume)[0].strip()
        dates = RE_DATE.findall(texte)
        date_pub = None
        if dates:
            j, m, y = dates[-1]
            date_pub = dt.datetime(int(y), MOIS[m.lower()], int(j))
        articles[lien] = dict(titre=titre, lien=lien, resume=resume,
                              tags=" ".join(tags), date_pub=date_pub)
    return list(articles.values())


def est_sanction(art):
    if "sanction" not in art["tags"].lower():
        return False
    return not any(x in art["titre"].lower() for x in TITRES_EXCLUS)


TAGS_IGNORES = {"sanction", "particulier", "professionnel"}


def _propre(t):
    return re.sub(r"\s+", " ", (t or "").replace("\u2019", "'")).strip()


def details_article(lien):
    """Lit la page de l'article CNIL et en tire, quand c'est possible :
    - la délibération (référence SAN + date), dans la rubrique « La délibération » en bas d'article ;
    - les manquements, à partir des intertitres « Un manquement à ... (article X du RGPD) » ;
    - l'entité (« concernant la société X ») et le montant (« amende de X euros ») ;
    - les mots clés thématiques que la CNIL attache à l'article (#Cookies, #Prospection...).
    """
    infos = dict(ref="", date=None, manquements="", entite="", montant="", mots_cles="", texte="", legifrance=[])
    try:
        soup = BeautifulSoup(telecharger(lien), "html.parser")
    except Exception as e:
        print(f"  ! impossible de lire {lien} : {e}")
        return infos
    corps = soup.find("main") or soup
    texte = _propre(corps.get_text(" ", strip=True))

    # 1) Délibération(s) : "Délibération SAN-2025-017 du 30 décembre 2025"
    #    ou "Délibération de la formation restreinte n° SAN – 2026-010 du 21 juillet 2026 concernant la société EXTIA"
    refs, dates = [], []
    for m in re.finditer(r"Délibération[^.]{0,80}?(SAN\s*[-–]\s*\d{4}\s*[-–]\s*\d{3})\s*(?:du\s+(\d{1,2}(?:er)?\s+\w+\s+\d{4}))?"
                         r"(?:[^-]{0,60}?(?:concernant|à l'encontre de)\s+(?:la société |les sociétés |l'|le |la )?([^-|]+?)(?=\s+-\s+Légifrance|\s+Texte\b|\s+Délibération|$))?",
                         texte):
        ref = re.sub(r"\s", "", m.group(1)).replace("–", "-")
        if ref not in refs:
            refs.append(ref)
            if m.group(2):
                dates.append(parse_date(m.group(2)))
            if m.group(3) and not infos["entite"] and len(m.group(3)) < 80 and not m.group(3).strip().upper().startswith("X"):
                infos["entite"] = m.group(3).strip()
    infos["ref"] = " et ".join(f"Délibération {r}" for r in refs)
    infos["date"] = next((d for d in dates if d), None)

    # 2) Manquements : intertitres "Un manquement à l'obligation de ... (article 6 du RGPD)"
    #    (variantes rencontrées : h2 ou h3, "article 32 RGPD", "article L.34-5 du CPCE", "articles 5-1-a) et 5-2 du RGPD")
    lignes, vus = [], set()
    for h in corps.find_all(["h2", "h3", "h4"]):
        t = _propre(h.get_text(" "))
        m = re.search(r"^(?:un |des |le |les )?(manquements?\b.*?)\s*\((?:articles?|art\.)\s+(.+?)\s+((?:du |de la )?(?:RGPD|règlement[^)]*|loi[^)]*|CPCE|code des postes[^)]*|LIL))\s*\)\s*$", t, re.I)
        if not m:
            continue
        desc, nums, src = m.group(1), m.group(2), m.group(3).lower()
        if "loi" in src or "lil" in src:
            nums, suffixe = nums, " LIL"
        elif "cpce" in src or "postes" in src:
            suffixe = " CPCE"
        else:
            suffixe = ""
        ligne = f"Art {nums}{suffixe} ({desc[0].lower() + desc[1:]})"
        if ligne not in vus:
            vus.add(ligne)
            lignes.append(ligne)
    #    Autre présentation : liste à puces en gras sous « Les manquements sanctionnés »,
    #    avec l'article cité dans la phrase d'introduction (cas fréquent pour les cookies).
    if not lignes:
        titre = next((h for h in corps.find_all(["h2", "h3"]) if "manquements sanctionn" in _propre(h.get_text()).lower()), None)
        if titre:
            bloc, puces = [], []
            for el in titre.find_all_next():
                if el.name == "h2" and el is not titre:
                    break
                if el.name in ("p", "h3"):
                    bloc.append(_propre(el.get_text(" ")))
                if el.name == "strong":
                    puces.append(_propre(el.get_text(" ")).rstrip(" :"))
            intro = " ".join(bloc)
            m = re.search(r"articles?\s+([\dL][\w.\-]*(?:\s+et\s+[\w.\-]+)?)\s+(du RGPD|de la loi Informatique et Libertés|du CPCE)", intro, re.I)
            if m:
                num, src = m.group(1), m.group(2)
            else:
                m = re.search(r"(RGPD|loi Informatique et Libertés|CPCE)\s*\((?:articles?|art\.)\s+([\w.\-]+)\)", intro, re.I)
                num, src = (m.group(2), m.group(1)) if m else (None, None)
            if num and puces:
                suffixe = " LIL" if "loi" in src.lower() else (" CPCE" if "CPCE" in src else "")
                lignes.append(f"Art {num}{suffixe} (" + " ; ".join(p[0].lower() + p[1:] for p in puces) + ")")
    infos["manquements"] = "\n".join(lignes)

    # 3) Montant : "amende de 3,5 millions d'euros" / "amende administrative d'un montant de 5 000 euros"
    m = re.search(r"amendes?(?: administratives?)?(?: d'un montant)? de ((?:respectivement )?\d[\d\s.,]*(?:\s*(?:et|,)\s*\d[\d\s.,]*)?\s*(?:millions?|milliards?)?)\s*(?:d'euros|euros|€)", texte, re.I)
    if m:
        infos["montant"] = m.group(1).replace("respectivement ", "").strip() + " €"

    # 4) Mots clés : tags thématiques de la CNIL + phrase d'accroche de l'article
    tags = []
    for a in corps.find_all("a", href=re.compile(r"/tag/")):
        t = _propre(a.get_text()).lstrip("#")
        if t and t.lower() not in TAGS_IGNORES and t not in tags:
            tags.append(t)
    infos["mots_cles"] = "\n".join(tags)
    infos["texte"] = texte

    # 5) Lien(s) vers la délibération publiée sur Légifrance
    for a in corps.find_all("a", href=re.compile(r"legifrance\.gouv\.fr/cnil/id/CNILTEXT", re.I)):
        url = a["href"].split("?")[0]
        if url not in infos["legifrance"]:
            infos["legifrance"].append(url)
    if re.search(r"sans qu'il soit[^.]*utile de nommer|ne peut pas divulguer le nom|anonymis", texte, re.I) and not infos["entite"]:
        infos["entite"] = "Identité de la société non disponible"
    return infos


def entite(titre):
    m = re.search(r"à l[’']encontre (?:de la société |des sociétés |de l[’']|du |de la |des |de |d[’'])(.+)$", titre, re.I)
    return m.group(1).strip() if m else ""


def montant(titre, resume):
    for t in (titre, resume):
        m = re.search(r"(\d[\d\s.,]*\s*(?:millions?|milliards?)?)\s*(?:d[’']euros|euros|€)", t, re.I)
        if m:
            return m.group(1).strip() + " €"
    return ""


def ajuster_mise_en_page(ws):
    """Rend l'Excel lisible dès l'ouverture : retour à la ligne, texte en haut des cellules,
    et hauteur de chaque ligne calculée d'après la longueur du texte (Excel ne le fait pas
    tout seul pour un fichier généré par programme)."""
    import math
    largeurs = {c: (ws.column_dimensions[get_column_letter(c)].width or 10) for c in range(1, COL_CLE + 1)}
    ws.freeze_panes = "A2"
    # La colonne H garde les adresses Légifrance (utile quand il y a plusieurs délibérations),
    # mais elle est masquée : on clique directement sur la référence en colonne 1.
    ws.column_dimensions[get_column_letter(COL_LEGI)].hidden = True
    ws.column_dimensions[get_column_letter(COL_CLE)].hidden = True
    for r in range(2, ws.max_row + 1):
        vide = all(ws.cell(r, c).value in (None, "") for c in range(1, 7))
        if vide:
            ws.row_dimensions[r].height = None
            continue
        nb_lignes = 1
        for c in range(1, COL_AJOUT + 1):
            cell = ws.cell(r, c)
            if c == COL_LEGI:
                continue
            al = copy(cell.alignment)
            al.wrap_text, al.vertical, al.shrink_to_fit = True, "top", False
            cell.alignment = al
            if c == 2 and isinstance(cell.value, dt.datetime):
                cell.number_format = "dd/mm/yyyy"
            texte = "" if cell.value is None else (f"{cell.value:%d/%m/%Y}" if isinstance(cell.value, dt.datetime) else str(cell.value))
            car_par_ligne = max(1, int(largeurs[c] * 1.15))
            n = sum(max(1, math.ceil(len(par.strip()) / car_par_ligne)) for par in texte.split("\n")) if texte else 1
            nb_lignes = max(nb_lignes, n)
        ws.row_dimensions[r].height = min(409, nb_lignes * 15 + 6)


# ---------------------------------------------------------------------------
# Tableau officiel « Les sanctions prononcées par la CNIL » (historique complet)
# ---------------------------------------------------------------------------
RE_JJMMAAAA = re.compile(r"\b(\d{1,2})[/.](\d{1,2})[/.](\d{4}|\d{2})\b")
RE_REF_SAN = re.compile(r"SAN\s*[-–]\s*\d{3,4}\s*[-–]\s*\d{3}", re.I)

# Mots clés déduits des manquements (pour la colonne C et la recherche)
THEMES = [
    ("Cookies et traceurs", r"cookie|traceur"),
    ("Prospection commerciale", r"prospection commerciale|L\.? ?34-5|démarchage"),
    ("Prospection politique", r"prospection politique|élection|candidat|parti politique"),
    ("Vidéosurveillance", r"vidéo"),
    ("Géolocalisation", r"géolocalisation"),
    ("Sécurité des données", r"sécurité"),
    ("Violation de données", r"violation de données|fuite"),
    ("Droits des personnes", r"droit d'accès|droit d.effacement|droit à l'effacement|droit d'opposition|rectification|exercice des droits|portabilité"),
    ("Information des personnes", r"information des personnes|transparence"),
    ("Durée de conservation", r"durée de conservation|conservation"),
    ("Minimisation", r"minimisation|pertinen|excessi"),
    ("Base légale", r"base légale|licite|licéité"),
    ("Données de santé", r"santé|médical|médecin|dentiste|clinique|hôpital|hospitali"),
    ("Sous-traitance", r"sous-traitan"),
    ("Coopération avec la CNIL", r"coopération|réponse à l'injonction|injonction|astreinte|mise en demeure"),
    ("Analyse d'impact", r"analyse d'impact"),
    ("Registre", r"registre"),
    ("Délégué à la protection des données", r"délégué à la protection"),
    ("Transferts hors UE", r"transfert"),
]


def _date_fr(txt):
    m = RE_JJMMAAAA.search(txt or "")
    if m:
        j, mo, a = (int(x) for x in m.groups())
        a += 2000 if a < 100 else 0
        try:
            return dt.datetime(a, mo, j)
        except ValueError:
            return None
    return parse_date(txt)


def _texte_cellule(td):
    """Texte d'une cellule en gardant les retours à la ligne (<br>, <p>, <li>),
    sans couper les mots en italique ou en gras."""
    td = copy(td)
    for br in td.find_all("br"):
        br.replace_with("\n")
    for bloc in td.find_all(["p", "li", "div"]):
        bloc.append("\n")
    lignes = []
    for l in td.get_text("").replace("\xa0", " ").replace("\u2019", "'").split("\n"):
        # dans certaines années, les manquements sont séparés par de longs blancs
        for morceau in re.split(r"\s{3,}(?=[A-ZÀ-Ý])", l):
            morceau = re.sub(r"\s+", " ", morceau).strip(" ;")
            morceau = re.sub(r"\s+([,)])", r"\1", morceau)
            if morceau:
                lignes.append(morceau)
    return lignes


def _id_legifrance(url):
    m = re.search(r"CNILTEXT\d+", url or "")
    return m.group(0) if m else None


def _url_legifrance(id_):
    return f"https://www.legifrance.gouv.fr/cnil/id/{id_}"


def _montant_num(txt):
    """Premier montant en euros trouvé dans le texte, en nombre (pour comparer)."""
    t = (txt or "").replace("\xa0", " ").replace("\u2019", "'")
    m = re.search(r"(\d[\d .,]*?)\s*(milliards?|millions?|M€)?\s*(?:d'euros|euros|€)", t, re.I)
    if not m:
        return None
    brut = m.group(1).replace(" ", "")
    if re.fullmatch(r"\d{1,3}(\.\d{3})+", brut):      # 300.000
        brut = brut.replace(".", "")
    brut = brut.replace(",", ".")
    try:
        v = float(brut)
    except ValueError:
        return None
    unite = (m.group(2) or "").lower()
    if unite.startswith("milliard"):
        v *= 1e9
    elif unite.startswith("million") or unite == "m€":
        v *= 1e6
    return round(v)


def _cle(date, organisme, decision):
    norm = lambda t: re.sub(r"[^a-z0-9]", "", htmlmod.unescape(t or "").lower())
    return f"{date:%Y%m%d}|{norm(organisme)[:60]}|{norm(decision)[:60]}"


def lire_tableau_officiel():
    """Lit le tableau officiel de la CNIL, toutes années confondues.
    Renvoie une liste de dict : date, organisme, manquements, decision, legifrance, ref, simplifiee, cle."""
    soup = BeautifulSoup(telecharger(URL_TABLEAU), "html.parser")
    corps = soup.find("main") or soup
    lignes, vus_cles = [], {}
    for table in corps.find_all("table"):
        pos = dict(date=0, org=1, manq=2, dec=3)
        entetes = [_propre(th.get_text(" ")).lower() for th in table.find_all("th")]
        for i, e in enumerate(entetes):
            if "date" in e:
                pos["date"] = i
            elif "organisme" in e or "nom" in e or "entit" in e:
                pos["org"] = i
            elif "manquement" in e or "thème" in e or "theme" in e:
                pos["manq"] = i
            elif "décision" in e or "decision" in e or "sanction" in e:
                pos["dec"] = i
        for tr in table.find_all("tr"):
            tds = tr.find_all("td")
            if len(tds) < 3:
                continue
            cell = lambda k: _texte_cellule(tds[pos[k]]) if pos[k] < len(tds) else []
            date = _date_fr(" ".join(cell("date"))) or next(
                (d for d in (_date_fr(" ".join(_texte_cellule(td))) for td in tds) if d), None)
            if not date:
                continue                                   # ligne d'en-tête ou vide
            organisme = " ".join(cell("org"))
            simplifiee = bool(re.search(r"proc[ée]dure simplifi", organisme, re.I))
            organisme = re.sub(r"\s*\(?\s*proc[ée]dure simplifi[ée]e\s*\)?", "", organisme, flags=re.I).strip()
            decision = [re.sub(r"\s*voir la d[ée]lib[ée]ration.*$", "", l, flags=re.I) for l in cell("dec")]
            decision = [l for l in decision if l]
            ids, ref = [], ""
            for a in tr.find_all("a", href=True):
                id_ = _id_legifrance(a["href"])
                if id_ and id_ not in ids:
                    ids.append(id_)
                if not ref:
                    m = RE_REF_SAN.search(a.get("title", "") + " " + a.get_text(" "))
                    if m:
                        ref = re.sub(r"\s", "", m.group(0)).replace("–", "-").upper()
                        ref = re.sub(r"^SAN-0(\d{2})-", r"SAN-20\1-", ref)   # coquille CNIL « SAN-026-001 »
                    else:
                        m = re.search(r"d[ée]lib[ée]ration\s+(?:n°\s*)?([\w\-–]+)\s+du\s+\d", a.get("title", ""), re.I)
                        if m and re.search(r"\d", m.group(1)):
                            ref = m.group(1).replace("–", "-")
            cle = _cle(date, organisme, " ".join(decision))
            vus_cles[cle] = vus_cles.get(cle, 0) + 1
            if vus_cles[cle] > 1:
                cle += f"#{vus_cles[cle]}"                 # deux lignes identiques le même jour
            lignes.append(dict(date=date, organisme=organisme, manquements=cell("manq"),
                               decision="\n".join(decision), legifrance=[_url_legifrance(i) for i in ids],
                               ref=ref, simplifiee=simplifiee, cle=cle))
    return lignes


def infos_legifrance(url):
    """Lit la délibération sur Légifrance pour retrouver le nom de l'organisme
    sanctionné (le tableau CNIL ne donne que son type d'activité) et sa référence."""
    soup = BeautifulSoup(telecharger(url), "html.parser")
    titre = _propre((soup.find("h1") or soup.title or soup).get_text(" "))
    texte = titre + " " + _propre((soup.find("main") or soup).get_text(" "))[:4000]
    ref = ""
    m = RE_REF_SAN.search(texte)
    if m:
        ref = re.sub(r"\s", "", m.group(0)).replace("–", "-").upper()
    else:                                                   # anciennes délibérations : « n° 2012-345 »
        m = re.search(r"Délibération[^0-9]{0,60}?(\d{4}-\d{2,4})\b", titre)
        ref = m.group(1) if m else ""
    nom = ""
    m = re.search(r"(?:à l'encontre (?:de la société|des sociétés|de l'|du|de la|des|de|d')|concernant (?:la société|les sociétés|l'|le|la|les))\s*"
                  r"(.{2,90}?)(?=\s+(?:La (?:Commission|formation)|Vu\b|Délibération|Texte|Après)|[,;]|$)", texte)
    if m:
        cand = m.group(1).strip(" .")
        lettres = [c for c in cand if c.isalpha()]
        majuscules = sum(c.isupper() for c in lettres) / max(1, len(lettres))
        if majuscules > 0.6 and not re.fullmatch(r"[XYZ]+\.?|[A-Z]\.?", cand):
            nom = cand
    return dict(ref=ref, nom=nom)


def mots_cles_depuis(manquements, simplifiee):
    texte = " ".join(manquements).lower().replace("\u2019", "'")
    themes = [nom for nom, motif in THEMES if re.search(motif, texte, re.I)]
    entete = LIBELLE_SIMPLIFIEE if simplifiee else ""
    return "\n".join(([entete] if entete else []) + themes)


def rattraper_historique(ws, dry_run=False):
    """Ajoute à la feuille toutes les sanctions du tableau officiel qui n'y sont pas encore."""
    try:
        officiel = lire_tableau_officiel()
    except Exception as e:
        print(f"! Tableau officiel des sanctions illisible : {e}")
        return 0
    if not officiel:
        print("! Aucune ligne lue dans le tableau officiel (structure de la page modifiée ?)")
        return 0
    annees = sorted({l["date"].year for l in officiel})
    print(f"Tableau officiel CNIL : {len(officiel)} sanctions lues ({annees[0]}–{annees[-1]})")

    ids_connus, cles_connues, date_montant = set(), set(), set()
    for r in range(2, ws.max_row + 1):
        for c in (1, COL_LEGI):
            v = ws.cell(r, c).hyperlink.target if c == 1 and ws.cell(r, c).hyperlink else ws.cell(r, c).value
            ids_connus.update(re.findall(r"CNILTEXT\d+", str(v or "")))
        if ws.cell(r, COL_CLE).value:
            cles_connues.add(str(ws.cell(r, COL_CLE).value))
        d, mt = ws.cell(r, 2).value, _montant_num(str(ws.cell(r, 6).value or ""))
        if isinstance(d, dt.datetime) and mt:
            date_montant.add((d.date(), mt))

    a_ajouter = []
    for l in officiel:
        ids = {_id_legifrance(u) for u in l["legifrance"]}
        if l["cle"] in cles_connues or (ids & ids_connus):
            continue
        mt = _montant_num(l["decision"])
        if not l["simplifiee"] and mt and (l["date"].date(), mt) in date_montant:
            continue                                       # déjà saisie à partir de l'article CNIL
        if ids & {i for x in a_ajouter for i in map(_id_legifrance, x["legifrance"])}:
            continue                                       # même délibération listée deux fois par la CNIL
        a_ajouter.append(l)
    print(f"  -> {len(a_ajouter)} sanction(s) à ajouter au tableau")
    if not a_ajouter or dry_run:
        return len(a_ajouter)

    # Nom réel de l'organisme et référence exacte, quand la délibération est publiée
    lues, echecs = 0, 0
    for l in sorted(a_ajouter, key=lambda x: x["date"], reverse=True):
        if not l["legifrance"] or lues >= MAX_LEGIFRANCE or echecs >= 5:
            continue
        try:
            info = infos_legifrance(l["legifrance"][0])
            lues += 1
            l["nom"] = info["nom"]
            l["ref"] = l["ref"] or info["ref"]
            time.sleep(0.5)
        except Exception as e:
            echecs += 1
            print(f"  ! Légifrance illisible ({l['legifrance'][0]}) : {e}")
    print(f"  {lues} délibération(s) lue(s) sur Légifrance")

    modele = 2
    derniere = max([r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value not in (None, "")] or [1])
    maintenant = dt.datetime.now().replace(microsecond=0)
    for l in sorted(a_ajouter, key=lambda x: x["date"]):
        if l["ref"]:
            ref = f"Délibération {l['ref']}"
        elif l["legifrance"]:
            ref = "Délibération publiée sur Légifrance"
        elif l["simplifiee"]:
            ref = "Procédure simplifiée (décision non publiée)"
        else:
            ref = "Délibération non publiée"
        organisme = l["organisme"].strip() or "Non précisé"
        norm = lambda t: re.sub(r"[^a-z0-9]", "", t.lower())
        if l.get("nom") and norm(organisme) not in norm(l["nom"]):
            entite_txt = f"{l['nom']}\n({organisme})"
        else:
            entite_txt = l.get("nom") or organisme
        ligne = [ref, l["date"], mots_cles_depuis(l["manquements"] + [organisme], l["simplifiee"]), entite_txt,
                 "\n".join(l["manquements"]), l["decision"] or "Non précisé", None,
                 "\n".join(l["legifrance"]) or None, maintenant, l["cle"]]
        derniere += 1
        for col, val in enumerate(ligne, 1):
            c = ws.cell(derniere, col, val)
            c._style = copy(ws.cell(modele, min(col, 6))._style)
            c.fill = PatternFill(fill_type=None)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if l["legifrance"]:
            ws.cell(derniere, 1).hyperlink = l["legifrance"][0]
            ws.cell(derniere, 1).font = Font(name=ws.cell(modele, 1).font.name, color="0563C1", underline="single")
        ws.cell(derniere, COL_AJOUT).number_format = "dd/mm/yyyy hh:mm"
    return len(a_ajouter)


def _ligne_existante(ws, ids):
    """Numéro de ligne déjà présente pour l'une de ces délibérations Légifrance (sinon None)."""
    if not ids:
        return None
    for r in range(2, ws.max_row + 1):
        v = str(ws.cell(r, COL_LEGI).value or "")
        if any(i in v for i in ids):
            return r
    return None


# ---------------------------------------------------------------------------
# Liens vers les articles « Actualités » de la CNIL pour TOUTES les lignes
# ---------------------------------------------------------------------------
URL_TAG_SANCTION = "https://www.cnil.fr/fr/mots-cles/sanction"
MAX_PAGES_ARCHIVE = 900          # pages de cnil.fr/fr/actualite parcourues lors de la reprise complète (une seule fois)
FEUILLE_VUS = "_articles_vus"    # feuille masquée de l'Excel : articles déjà examinés (évite de les relire chaque jour)
MARQUEUR_ARCHIVE = "__archive_parcourue__"


def _norm(t):
    t = htmlmod.unescape(str(t or "")).lower()
    for a, b in (("éèêë", "e"), ("àâä", "a"), ("îï", "i"), ("ôö", "o"), ("ùûü", "u"), ("ç", "c")):
        for x in a:
            t = t.replace(x, b)
    return re.sub(r"[^a-z0-9]", "", t)


def _feuille_vus(wb):
    if FEUILLE_VUS in wb.sheetnames:
        fv = wb[FEUILLE_VUS]
    else:
        fv = wb.create_sheet(FEUILLE_VUS)
        fv.append(["Article CNIL", "Résultat", "Examiné le"])
    fv.sheet_state = "hidden"
    vus = {str(r[0]) for r in fv.iter_rows(min_row=2, values_only=True) if r and r[0]}
    return fv, vus


def est_article_sanction(art):
    t = art["titre"].lower()
    if any(x in t for x in TITRES_EXCLUS):
        return False
    return est_sanction(art) or bool(re.search(r"\bsanction|\bamende|avertissement public|sanctionn", t))


def lister_articles(base, pages_max, pause=0.3):
    """Parcourt une liste paginée de cnil.fr et renvoie les articles de sanction trouvés."""
    trouves, deja_vus = {}, set()
    for p in range(pages_max):
        try:
            arts = extraire_articles(telecharger(f"{base}?page={p}"), base)
        except Exception as e:
            print(f"  ! page {p} de {base} illisible : {e}")
            break
        liens = {a["lien"] for a in arts}
        if not arts or liens <= deja_vus:           # fin de la pagination
            break
        deja_vus |= liens
        for a in arts:
            if est_article_sanction(a):
                trouves.setdefault(a["lien"], a)
        if p and p % 50 == 0:
            print(f"  … {p} pages lues, {len(trouves)} articles de sanction repérés")
        time.sleep(pause)
    return list(trouves.values())


def trouver_ligne(ws, ids, date, montant_eur, nom):
    """Ligne du tableau qui correspond à un article : même délibération Légifrance,
    sinon même date et même montant, sinon même nom d'organisme à quelques mois près."""
    r = _ligne_existante(ws, ids)
    if r:
        return r
    lignes = [(r, ws.cell(r, 2).value) for r in range(2, ws.max_row + 1) if isinstance(ws.cell(r, 2).value, dt.datetime)]
    if date and montant_eur:
        for r, d in lignes:
            if abs((d - date).days) <= 3 and _montant_num(str(ws.cell(r, 6).value or "")) == montant_eur:
                return r
    n = _norm(nom)
    if date and len(n) >= 4 and n not in {"societe", "lasociete"}:
        proches = [(abs((d - date).days), r) for r, d in lignes
                   if abs((d - date).days) <= 200 and n in _norm(ws.cell(r, 4).value)]
        if proches:
            return min(proches)[1]
    return None


def _mettre_lien(ws, r, url):
    c = ws.cell(r, COL_LIEN, url)
    c.hyperlink = url
    c.font = Font(name=ws.cell(r, 1).font.name, color="0563C1", underline="single")


def relier_articles(wb, ws, dry_run=False):
    """Ajoute le lien vers l'article « Actualités » de la CNIL aux lignes qui n'en ont pas
    (sanctions reprises du tableau officiel). La première fois, toute l'archive des actualités
    est parcourue ; ensuite, seules les pages récentes."""
    fv, vus = _feuille_vus(wb)
    complet = MARQUEUR_ARCHIVE in vus
    if complet:
        articles = lister_articles(URL_TAG_SANCTION, 2)
    else:
        print("Reprise des liens vers les articles CNIL : parcours de toute l'archive des actualités (une seule fois)…")
        articles = lister_articles(URL_ACTUS, MAX_PAGES_ARCHIVE) + lister_articles(URL_TAG_SANCTION, 50)
        articles = list(dict((a["lien"], a) for a in articles).values())
    dans_tableau = {str(ws.cell(r, COL_LIEN).value).strip() for r in range(2, ws.max_row + 1) if ws.cell(r, COL_LIEN).value}
    a_voir = [a for a in articles if a["lien"] not in vus and a["lien"] not in dans_tableau]
    print(f"Articles de sanction à rattacher : {len(a_voir)} (sur {len(articles)} repérés)")
    if dry_run:
        return 0

    maintenant = dt.datetime.now().replace(microsecond=0)
    recaps, relies = [], 0
    for art in sorted(a_voir, key=lambda a: a["date_pub"] or dt.datetime.min):
        # Récapitulatifs des sanctions en procédure simplifiée : rattachés plus bas, par période
        if re.search(r"simplifi", art["titre"], re.I):
            if art["date_pub"]:
                recaps.append(art)
            fv.append([art["lien"], "récapitulatif procédure simplifiée", maintenant])
            continue
        d = details_article(art["lien"])
        time.sleep(0.3)
        ids = [_id_legifrance(u) for u in d["legifrance"]]
        date = d["date"] or parse_date(art["resume"]) or art["date_pub"]
        mt = _montant_num(montant(art["titre"], art["resume"]) or d["montant"])
        nom = entite(art["titre"]) or (d["entite"] if not d["entite"].startswith("Identité") else "")
        if not re.search(r"[A-ZÀ-Ý]{2,}", nom):          # « un vendeur de mobilier » n'est pas un nom de société
            nom = ""
        r = trouver_ligne(ws, [i for i in ids if i], date, mt, nom)
        if not r:
            fv.append([art["lien"], "aucune ligne correspondante", maintenant])
            continue
        if ws.cell(r, COL_LIEN).value in (None, ""):
            _mettre_lien(ws, r, art["lien"])
            # on profite de l'article pour préciser la ligne (référence, entité, articles du RGPD)
            if d["ref"] and "SAN" not in str(ws.cell(r, 1).value or ""):
                ws.cell(r, 1, d["ref"])
            if nom and _norm(nom) not in _norm(ws.cell(r, 4).value):
                ws.cell(r, 4, f"{nom}\n({ws.cell(r, 4).value})" if ws.cell(r, 4).value else nom)
            if d["manquements"]:
                ws.cell(r, 5, d["manquements"])
            if d["legifrance"] and not ws.cell(r, COL_LEGI).value:
                ws.cell(r, COL_LEGI, "\n".join(d["legifrance"]))
                ws.cell(r, 1).hyperlink = d["legifrance"][0]
            relies += 1
        fv.append([art["lien"], f"rattaché à la ligne {r}", maintenant])

    # Procédure simplifiée : la CNIL publie un article récapitulatif par période.
    # Chaque sanction simplifiée reçoit le lien du premier récapitulatif publié après elle.
    if recaps:
        recaps.sort(key=lambda a: a["date_pub"])
        for r in range(2, ws.max_row + 1):
            d = ws.cell(r, 2).value
            simpl = LIBELLE_SIMPLIFIEE.lower() in (str(ws.cell(r, 1).value) + str(ws.cell(r, 3).value)).lower()
            if not simpl or not isinstance(d, dt.datetime) or ws.cell(r, COL_LIEN).value:
                continue
            rec = next((a for a in recaps if a["date_pub"] >= d and (a["date_pub"] - d).days <= 400), None)
            if rec:
                _mettre_lien(ws, r, rec["lien"])
                relies += 1

    if not complet:
        fv.append([MARQUEUR_ARCHIVE, "archive des actualités parcourue", maintenant])
    print(f"  -> {relies} ligne(s) reliée(s) à un article CNIL")
    return relies


def generer_page():
    """Construit site/index.html (le tableau que voient les boss) à partir de l'Excel."""
    ws = load_workbook(FICHIER_EXCEL, data_only=True).worksheets[0]
    entetes = [str(ws.cell(1, c).value or "").strip() for c in range(1, 7)]
    lignes = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, COL_LEGI + 1)]
        if all(v in (None, "") for v in vals[:6]):
            continue
        lignes.append(vals)
    lignes.sort(key=lambda v: v[1] if isinstance(v[1], dt.datetime) else dt.datetime.min, reverse=True)

    def cellule(v):
        if isinstance(v, dt.datetime):
            return f'<td data-sort="{v:%Y%m%d}">{v:%d/%m/%Y}</td>'
        t = htmlmod.escape(str(v or "").strip()).replace("\n", "<br>")
        cls = ' class="todo"' if "À compléter" in t else ""
        return f"<td{cls}>{t}</td>"

    def cellule_ref(ref, liens):
        """Colonne 1 : la référence de la délibération devient un lien vers Légifrance.
        S'il y a plusieurs délibérations (ex. Free), chacune a son propre lien."""
        texte = str(ref or "").strip()
        urls = re.findall(r"https?://\S+", str(liens or ""))
        if not urls:
            return cellule(texte)
        lien = lambda u, t: f'<a href="{htmlmod.escape(u)}" target="_blank" rel="noopener" title="Voir la délibération sur Légifrance">{t}</a>'
        morceaux = [m.strip() for m in re.split(r"\s+et\s+(?=Délibération)", texte)]
        if len(urls) > 1 and len(morceaux) == len(urls):
            html = " et<br>".join(lien(u, htmlmod.escape(m)) for u, m in zip(urls, morceaux))
        else:
            html = lien(urls[0], htmlmod.escape(texte).replace("\n", "<br>"))
        return f"<td>{html}</td>"

    corps, annees, nb_simpl = "", set(), 0
    for v in lignes:
        lien = str(v[6] or "").strip()
        a = f'<a href="{htmlmod.escape(lien)}" target="_blank" rel="noopener">Voir</a>' if lien.startswith("http") else ""
        annee = v[1].year if isinstance(v[1], dt.datetime) else ""
        annees.add(annee) if annee else None
        simpl = LIBELLE_SIMPLIFIEE.lower() in (str(v[0] or "") + " " + str(v[2] or "")).lower()
        nb_simpl += simpl
        attrs = f' data-annee="{annee}" data-type="{"simplifiee" if simpl else "ordinaire"}"'
        corps += f"<tr{attrs}>" + cellule_ref(v[0], v[7]) + "".join(cellule(x) for x in v[1:6]) + f"<td>{a}</td></tr>\n"
    maj = dt.datetime.now(ZoneInfo("Europe/Paris"))
    th = "".join(f"<th>{htmlmod.escape(e)}</th>" for e in entetes) + "<th>Article CNIL</th>"
    options_annees = "".join(f'<option value="{y}">{y}</option>' for y in sorted(annees, reverse=True))
    page = PAGE_HTML.format(th=th, corps=corps, n=len(lignes), maj=f"{maj:%d/%m/%Y à %H:%M}",
                            annees=options_annees, n_simpl=nb_simpl, n_ord=len(lignes) - nb_simpl,
                            periode=f"{min(annees)}–{max(annees)}" if annees else "")
    DOSSIER_SITE.mkdir(exist_ok=True)
    (DOSSIER_SITE / "index.html").write_text(page, encoding="utf-8")
    wb_dl = load_workbook(FICHIER_EXCEL)
    ajuster_mise_en_page(wb_dl.worksheets[0])
    wb_dl.save(DOSSIER_SITE / FICHIER_EXCEL.name)
    print(f"Page générée : {DOSSIER_SITE / 'index.html'} ({len(lignes)} sanctions)")


PAGE_HTML = """<!doctype html>
<html lang="fr"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Veille des sanctions CNIL</title>
<style>
 body{{font-family:Arial,Helvetica,sans-serif;margin:0;background:#f5f6f8;color:#1d1d1f}}
 header{{background:#1f3a5f;color:#fff;padding:20px 28px}}
 header h1{{margin:0 0 6px;font-size:22px}} header p{{margin:0;opacity:.85;font-size:14px}}
 main{{padding:20px 28px}}
 .barre{{display:flex;gap:12px;flex-wrap:wrap;align-items:center;margin-bottom:14px}}
 input,select{{padding:8px 10px;border:1px solid #c7ccd4;border-radius:6px;font-size:14px;background:#fff}}
 input{{min-width:260px}} .compte{{font-size:13px;color:#6b7280}}
 .btn{{background:#1f3a5f;color:#fff;text-decoration:none;padding:8px 14px;border-radius:6px;font-size:14px}}
 .tab{{overflow-x:auto;background:#fff;border-radius:8px;box-shadow:0 1px 3px rgba(0,0,0,.08)}}
 table{{border-collapse:collapse;width:100%;font-size:13px}}
 th{{background:#e9edf3;text-align:left;padding:10px;position:sticky;top:0;cursor:pointer;white-space:nowrap}}
 td{{padding:10px;border-top:1px solid #eceef2;vertical-align:top;min-width:90px}}
 td:nth-child(3),td:nth-child(5){{min-width:260px}}
 tr:hover td{{background:#fafbfd}} .todo{{color:#b26a00;font-style:italic}}
 a{{color:#0563c1}} footer{{padding:10px 28px 30px;font-size:12px;color:#6b7280}}
</style></head><body>
<header><h1>Veille des sanctions CNIL</h1>
<p>{n} sanctions ({periode}) · mis à jour automatiquement chaque jour ouvré · dernière vérification le {maj}</p></header>
<main>
<div class="barre"><input id="q" placeholder="Rechercher (entité, article, mot clé…)" oninput="filtrer()">
<select id="type" onchange="filtrer()"><option value="">Toutes les procédures</option>
<option value="ordinaire">Procédure ordinaire ({n_ord})</option><option value="simplifiee">Procédure simplifiée ({n_simpl})</option></select>
<select id="annee" onchange="filtrer()"><option value="">Toutes les années</option>{annees}</select>
<span id="compte" class="compte"></span>
<a class="btn" href="Sanctions_CNIL.xlsx" download>Télécharger l'Excel</a></div>
<div class="tab"><table id="t"><thead><tr>{th}</tr></thead><tbody>
{corps}</tbody></table></div></main>
<footer>Sources : cnil.fr/fr/actualite (articles détaillés) et cnil.fr/fr/les-sanctions-prononcees-par-la-cnil (liste officielle de toutes les sanctions depuis 2011, y compris en procédure simplifiée, dont les décisions ne sont pas publiées). Les informations marquées « À compléter » ou « à vérifier » ont été extraites automatiquement.</footer>
<script>
function filtrer(){{var q=document.getElementById('q').value.toLowerCase(),ty=document.getElementById('type').value,an=document.getElementById('annee').value,n=0;
 document.querySelectorAll('#t tbody tr').forEach(function(r){{var ok=r.innerText.toLowerCase().indexOf(q)>-1&&(!ty||r.dataset.type===ty)&&(!an||r.dataset.annee===an);
 r.style.display=ok?'':'none';n+=ok}});document.getElementById('compte').textContent=n+' sanction'+(n>1?'s':'')+' affichée'+(n>1?'s':'');}}
filtrer();
document.querySelectorAll('#t th').forEach(function(th,i){{var asc=false;th.onclick=function(){{asc=!asc;
 var b=document.querySelector('#t tbody');var rs=Array.from(b.rows);
 rs.sort(function(x,y){{var a=x.cells[i].dataset.sort||x.cells[i].innerText,c=y.cells[i].dataset.sort||y.cells[i].innerText;
 return (asc?1:-1)*a.localeCompare(c,'fr',{{numeric:true}})}});rs.forEach(function(r){{b.appendChild(r)}});}}}});
</script></body></html>"""


def main(dry_run=False):
    wb = load_workbook(FICHIER_EXCEL)
    ws = wb.worksheets[0]
    if any(ws.cell(1, c).value in (None, "") for c in (COL_LIEN, COL_LEGI, COL_AJOUT, COL_CLE)):
        ws.cell(1, COL_LIEN, "Lien article CNIL")
        ws.cell(1, COL_LEGI, "Lien délibération (Légifrance)")
        ws.cell(1, COL_AJOUT, "Ajouté automatiquement le")
        ws.cell(1, COL_CLE, "Clé tableau CNIL (technique)")
        for c in (COL_LIEN, COL_LEGI, COL_AJOUT, COL_CLE):
            ws.cell(1, c)._style = copy(ws.cell(1, 1)._style)
    deja = {str(ws.cell(r, COL_LIEN).value).strip() for r in range(2, ws.max_row + 1) if ws.cell(r, COL_LIEN).value}

    trouves, nb_articles = [], 0
    for p in range(NB_PAGES):
        html = telecharger(f"{URL_ACTUS}?page={p}")
        arts = extraire_articles(html)
        nb_articles += len(arts)
        print(f"Page {p + 1} : {len(arts)} articles lus")
        trouves += [a for a in arts if est_sanction(a)]

    # Garde-fou : si aucun article n'est lu, la page de la CNIL a sans doute changé.
    # On fait échouer le robot (croix rouge + e-mail automatique de GitHub) au lieu de se taire.
    if nb_articles == 0:
        sys.exit("ERREUR : aucun article lu sur cnil.fr/fr/actualite. La structure du site a peut-être changé.")

    trouves = list(dict((a["lien"], a) for a in trouves).values())
    print(f"\nSanctions repérées sur cnil.fr ({len(trouves)}) :")
    for a in trouves:
        etat = "déjà dans le tableau" if a["lien"] in deja else "NOUVELLE"
        print(f"  - [{etat}] {a['titre']}")
    print()

    nouveaux = [a for a in trouves if a["lien"] not in deja]
    nouveaux.sort(key=lambda a: a["date_pub"] or dt.datetime.min)

    derniere = max(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value not in (None, ""))
    modele = derniere
    ajoutes = []
    for art in nouveaux:
        # On ne garde que les sanctions dont la délibération est publiée (référence SAN-xxxx dans l'article).
        # Les récapitulatifs de sanctions en procédure simplifiée (anonymes, sans délibération) sont ignorés.
        if re.search(r"simplifi", art["titre"] + " " + art["resume"], re.I):
            print(f"  (ignorée, procédure simplifiée) {art['titre']}")
            continue
        d = details_article(art["lien"])
        if not d["ref"]:
            print(f"  (ignorée, pas de délibération publiée) {art['titre']}")
            continue
        ajoutes.append(art)
        date_decision = d["date"] or parse_date(art["resume"]) or art["date_pub"]
        chapo = art["resume"].rstrip("…. ")
        m = re.search(re.escape(_propre(chapo)[:60]) + r"[^.]*\.", d["texte"]) if chapo else None
        if m:
            chapo = m.group(0)
        mots = chapo + ("\n" + d["mots_cles"] if d["mots_cles"] else "")
        ligne = [
            d["ref"] or "À compléter",
            date_decision,
            mots,
            entite(art["titre"]) or d["entite"] or "À compléter",
            d["manquements"] or "À compléter",
            montant(art["titre"], art["resume"]) or d["montant"] or "À compléter",
            art["lien"],
            "\n".join(d["legifrance"]) or "À compléter",
            dt.datetime.now().replace(microsecond=0),
        ]
        print(f"+ {art['titre']}")
        if dry_run:
            continue
        # La sanction figure peut-être déjà (reprise depuis le tableau officiel) : on complète cette ligne.
        existante = _ligne_existante(ws, [_id_legifrance(u) for u in d["legifrance"]])
        if existante:
            for col, val in enumerate(ligne, 1):
                if val not in (None, "", "À compléter"):
                    ws.cell(existante, col, val)
            ws.cell(existante, COL_LIEN).hyperlink = art["lien"]
            ws.cell(existante, COL_LIEN).font = Font(name=ws.cell(modele, 1).font.name, color="0563C1", underline="single")
            continue
        derniere += 1
        for col, val in enumerate(ligne, 1):
            c = ws.cell(derniere, col, val)
            c._style = copy(ws.cell(modele, min(col, 6))._style)
            c.fill = SURLIGNER_NOUVEAU
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(derniere, COL_LIEN).hyperlink = art["lien"]
        ws.cell(derniere, COL_LIEN).font = Font(name=ws.cell(modele, 1).font.name, color="0563C1", underline="single")
        if d["legifrance"]:
            ws.cell(derniere, 1).hyperlink = d["legifrance"][0]
            ws.cell(derniere, 1).font = Font(name=ws.cell(modele, 1).font.name, color="0563C1", underline="single")
        ws.cell(derniere, COL_AJOUT).number_format = "dd/mm/yyyy hh:mm"

    if not nouveaux:
        print("Aucun nouvel article de sanction.")

    # Historique complet : tableau officiel de toutes les sanctions (y compris procédure simplifiée)
    print()
    nb_hist = rattraper_historique(ws, dry_run=dry_run)

    # Lien vers l'article CNIL (Actualités) pour les sanctions qui n'en ont pas encore
    print()
    nb_liens = relier_articles(wb, ws, dry_run=dry_run)

    if not dry_run:
        if ajoutes or nb_hist or nb_liens or FEUILLE_VUS in wb.sheetnames:
            ajuster_mise_en_page(ws)
            wb.save(FICHIER_EXCEL)
        print(f"{len(ajoutes)} article(s), {nb_hist} sanction(s) du tableau officiel et {nb_liens} lien(s) d'article ajoutés dans {FICHIER_EXCEL.name}")
        generer_page()
    return ajoutes


if __name__ == "__main__":
    if "--page-seule" in sys.argv:
        generer_page()
    else:
        main(dry_run="--dry-run" in sys.argv)
